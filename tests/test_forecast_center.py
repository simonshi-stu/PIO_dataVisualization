from __future__ import annotations

import unittest
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from pio_platform.forecast_center import (
    build_forecast_center,
    build_part_planning_records,
)
from pio_platform.sop_workbook import build_sop_workbook_bytes


class ForecastCenterTests(unittest.TestCase):
    def setUp(self) -> None:
        months = pd.date_range("2023-01-01", periods=43, freq="MS")
        rows = []
        for index, month in enumerate(months):
            partial = 0.55 if month == months[-1] else 1.0
            rows.extend(
                [
                    {
                        "month": month.strftime("%Y-%m"),
                        "brand": "H",
                        "entityKey": "H::ELANTRA",
                        "modelName": "Elantra",
                        "plc": "Floor Mat",
                        "partNumber": "H-MAT",
                        "partDescription": "Carpet Floor Mat",
                        "lifecycleStatus": "Active",
                        "installationQuantity": (100 + index) * partial,
                        "pioRevenue": (100 + index) * 80 * partial,
                    },
                    {
                        "month": month.strftime("%Y-%m"),
                        "brand": "H",
                        "entityKey": "H::ELANTRA",
                        "modelName": "Elantra",
                        "plc": "Cargo Tray",
                        "partNumber": "H-TRAY",
                        "partDescription": "Cargo Tray",
                        "lifecycleStatus": "Active",
                        "installationQuantity": (40 + index / 2) * partial,
                        "pioRevenue": (40 + index / 2) * 110 * partial,
                    },
                    {
                        "month": month.strftime("%Y-%m"),
                        "brand": "K",
                        "entityKey": "K::SPORTAGE",
                        "modelName": "Sportage",
                        "plc": "Floor Mat",
                        "partNumber": "K-MAT",
                        "partDescription": "Carpet Floor Mat",
                        "lifecycleStatus": "Active",
                        "installationQuantity": (70 + index / 3) * partial,
                        "pioRevenue": (70 + index / 3) * 75 * partial,
                    },
                ]
            )
        self.facts = pd.DataFrame(rows)
        self.working_days = pd.DataFrame(
            {
                "month": months.strftime("%Y-%m"),
                "workingDays": [20 + (index % 3) for index in range(len(months))],
            }
        )
        self.wholesale = pd.DataFrame(
            {
                "month": list(months.strftime("%Y-%m")) * 3,
                "brand": ["HMA"] * len(months) + ["GMA"] * len(months) + ["KUS"] * len(months),
                "modelName": ["Elantra"] * len(months) + ["GV70"] * len(months) + ["Sportage"] * len(months),
                "modelKey": ["ELANTRA"] * len(months) + ["GV70"] * len(months) + ["SPORTAGE"] * len(months),
                "wholesaleUnits": [200 + index for index in range(len(months))]
                + [50 + index for index in range(len(months))]
                + [150 + index for index in range(len(months))],
            }
        )

    def test_revenue_nowcast_and_hierarchy_reconcile(self) -> None:
        payload = build_forecast_center(
            self.facts,
            self.working_days,
            self.wholesale,
            metric="revenue",
            level="model_plc",
            horizon=3,
            latest_sales_month_is_complete=False,
            latest_sales_date=pd.Timestamp("2026-07-22"),
            include_all_records=True,
        )
        self.assertEqual(payload["summary"]["nowcastMonths"], ["2026-07"])
        self.assertEqual(payload["summary"]["pureForecastMonths"], ["2026-08", "2026-09"])
        self.assertEqual(payload["summary"]["reconciliation"]["status"], "PASS")
        self.assertLessEqual(len(payload["topAccessories"]), 10)
        self.assertTrue(all(record["plc"] in {"Floor Mat", "Cargo Tray"} for record in payload["records"]))
        self.assertEqual({record["brand"] for record in payload["brandRecords"]}, {"HMA", "KUS"})
        self.assertEqual(
            payload["summary"]["accuracyScope"]["evaluatedGrain"],
            "month × selected official brand anchor(s) from HMA/GMA/KUS",
        )
        self.assertIn("day-of-month cutoff is not a working-day count", payload["summary"]["periodExplanation"])
        self.assertIn(
            "current business rule",
            payload["summary"]["anchorPolicy"]["denominatorPolicy"]["HMA"],
        )
        unit_revenue = {
            (record["modelName"], record["plc"]): record["expectedUnitRevenue"]
            for record in payload["modelPlcRecords"]
            if record["allocationRoute"] == "regular_allocation"
        }
        self.assertAlmostEqual(unit_revenue[("Elantra", "Floor Mat")], 80.0, places=6)
        self.assertAlmostEqual(unit_revenue[("Elantra", "Cargo Tray")], 110.0, places=6)
        self.assertAlmostEqual(unit_revenue[("Sportage", "Floor Mat")], 75.0, places=6)
        formulas = {item["name"]: item for item in payload["summary"]["formulaCatalog"]}
        self.assertIn("SumOfPIS_CRP_CFM_PRI", formulas["Monthly target actual"]["formula"])
        self.assertIn("recent-6", formulas["Exact-part unit revenue"]["formula"])

    def test_wholesale_keeps_hma_gma_and_kus_as_separate_anchors(self) -> None:
        payload = build_forecast_center(
            self.facts,
            self.working_days,
            self.wholesale,
            metric="wholesale_quantity",
            level="model",
            horizon=2,
            latest_sales_month_is_complete=False,
            latest_sales_date=pd.Timestamp("2026-07-22"),
        )
        self.assertEqual(
            {record["brand"] for record in payload["brandRecords"]},
            {"HMA", "GMA", "KUS"},
        )
        self.assertEqual(payload["summary"]["reconciliation"]["status"], "PASS")
        self.assertEqual(payload["summary"]["forecastMonths"][0], "2026-07")
        self.assertEqual(payload["summary"]["nowcastMonths"], ["2026-07"])

    def test_wholesale_models_stopped_before_current_year_are_zero(self) -> None:
        stopped = pd.DataFrame(
            {
                "month": ["2023-01", "2025-07", "2023-01", "2025-10"],
                "brand": ["KUS", "KUS", "HMA", "HMA"],
                "modelName": ["Stinger", "Stinger", "Nexo", "Nexo"],
                "modelKey": ["STINGER", "STINGER", "NEXO", "NEXO"],
                "wholesaleUnits": [40.0, 30.0, 25.0, 20.0],
            }
        )
        wholesale = pd.concat([self.wholesale, stopped], ignore_index=True)

        payload = build_forecast_center(
            self.facts,
            self.working_days,
            wholesale,
            metric="wholesale_quantity",
            level="model",
            horizon=2,
            latest_sales_month_is_complete=False,
            latest_sales_date=pd.Timestamp("2026-07-22"),
            include_all_records=True,
        )
        stopped_records = {
            record["modelName"]: record
            for record in payload["modelRecords"]
            if record["modelName"] in {"Nexo", "Stinger"}
        }

        self.assertEqual(stopped_records["Nexo"]["allocationRoute"], "excluded_lifecycle")
        self.assertEqual(stopped_records["Stinger"]["allocationRoute"], "excluded_lifecycle")
        self.assertTrue(
            all(
                forecast["value"] == 0
                for record in stopped_records.values()
                for forecast in record["forecast"]
            )
        )
        self.assertEqual(payload["summary"]["reconciliation"]["status"], "PASS")

    def test_reference_portfolio_is_revenue_only_and_brand_specific(self) -> None:
        payload = build_forecast_center(
            self.facts,
            self.working_days,
            self.wholesale,
            metric="revenue",
            level="brand",
            horizon=2,
            model_strategy="reference_portfolio",
            latest_sales_month_is_complete=False,
            latest_sales_date=pd.Timestamp("2026-07-22"),
        )
        methods = {
            record["brand"]: record["selectedModel"]
            for record in payload["brandRecords"]
        }

        self.assertTrue(methods["HMA"].startswith("hw_add_add__heuristic"))
        self.assertEqual(methods["KUS"], "working_day_adjusted_seasonal")
        self.assertEqual(
            payload["summary"]["modelGovernance"]["requestedStrategy"],
            "reference_portfolio",
        )
        self.assertEqual(
            payload["summary"]["modelGovernance"]["referenceMethodStatus"],
            "validated_implementation_applied_to_unvalidated_source",
        )
        with self.assertRaisesRegex(ValueError, "only for Revenue"):
            build_forecast_center(
                self.facts,
                self.working_days,
                self.wholesale,
                metric="quantity",
                level="brand",
                model_strategy="reference_portfolio",
            )

    def test_exact_parts_allocate_to_plc_parent(self) -> None:
        payload = build_forecast_center(
            self.facts,
            self.working_days,
            self.wholesale,
            metric="quantity",
            level="model_plc",
            horizon=2,
            latest_sales_month_is_complete=False,
            latest_sales_date=pd.Timestamp("2026-07-22"),
            include_all_records=True,
        )
        part_records = build_part_planning_records(
            self.facts,
            payload["modelPlcRecords"],
            metric="quantity",
            latest_complete_month=payload["summary"]["latestCompleteMonth"],
        )
        self.assertGreater(len(part_records), 0)
        parent = {
            (record["brand"], record["entityKey"], record["plc"], item["month"]): item["value"]
            for record in payload["modelPlcRecords"]
            for item in record["forecast"]
        }
        children: dict[tuple[str, str, str, str], float] = {}
        for record in part_records:
            key = (record["brand"], record["entityKey"], record["plc"], record["month"])
            children[key] = children.get(key, 0.0) + record["value"]
        for key, value in parent.items():
            self.assertAlmostEqual(children[key], value, places=6)

    def test_low_volume_and_new_models_are_routed_without_breaking_reconciliation(self) -> None:
        facts = self.facts.copy()
        low = facts[facts["modelName"] == "Elantra"].copy()
        low["entityKey"] = "H::LOW VOLUME"
        low["modelName"] = "Low Volume"
        low["installationQuantity"] = 0.1
        low["pioRevenue"] = 8.0
        low["lifecycleStatus"] = "Active"
        low["lifecycleStatusCode"] = "active"
        new = facts[facts["modelName"] == "Elantra"].tail(8).copy()
        new["entityKey"] = "H::NEW EV"
        new["modelName"] = "New EV"
        new["installationQuantity"] = 30.0
        new["pioRevenue"] = 2400.0
        new["lifecycleStatus"] = "New / limited history"
        new["lifecycleStatusCode"] = "new"
        routed = pd.concat([facts, low, new], ignore_index=True)

        payload = build_forecast_center(
            routed,
            self.working_days,
            self.wholesale,
            metric="quantity",
            level="model",
            horizon=2,
            latest_sales_month_is_complete=False,
            latest_sales_date=pd.Timestamp("2026-07-22"),
            include_all_records=True,
        )
        routes = {record["modelName"]: record["allocationRoute"] for record in payload["modelRecords"]}
        self.assertEqual(routes["Low Volume"], "excluded_low_volume")
        self.assertEqual(routes["New EV"], "new_model_proxy")
        self.assertEqual(payload["summary"]["reconciliation"]["status"], "PASS")

    def test_sop_workbook_has_governed_sheets_and_formulas(self) -> None:
        payloads = {}
        for metric in ["revenue", "quantity", "wholesale_quantity"]:
            payloads[metric] = build_forecast_center(
                self.facts,
                self.working_days,
                self.wholesale,
                metric=metric,
                level="brand",
                horizon=2,
                latest_sales_month_is_complete=False,
                latest_sales_date=pd.Timestamp("2026-07-22"),
                include_all_records=True,
            )
        part_quantity = build_part_planning_records(
            self.facts,
            payloads["quantity"]["modelPlcRecords"],
            metric="quantity",
            latest_complete_month=payloads["quantity"]["summary"]["latestCompleteMonth"],
        )
        part_revenue = build_part_planning_records(
            self.facts,
            payloads["revenue"]["modelPlcRecords"],
            metric="revenue",
            latest_complete_month=payloads["revenue"]["summary"]["latestCompleteMonth"],
        )
        output = build_sop_workbook_bytes(
            source_filename="source.xlsx",
            revenue=payloads["revenue"],
            quantity=payloads["quantity"],
            wholesale=payloads["wholesale_quantity"],
            part_quantity=part_quantity,
            part_revenue=part_revenue,
            working_days=self.working_days.to_dict("records"),
        )
        workbook = load_workbook(BytesIO(output), data_only=False)
        self.assertEqual(
            workbook.sheetnames,
            [
                "Executive_Summary",
                "Revenue_Forecast",
                "Quantity_Forecast",
                "Part_Planning",
                "Wholesale_Drivers",
                "Model_Performance",
                "QA_Assumptions",
            ],
        )
        self.assertTrue(str(workbook["QA_Assumptions"]["B5"].value).startswith("="))
        self.assertTrue(str(workbook["Part_Planning"]["O6"].value).startswith("="))


if __name__ == "__main__":
    unittest.main()
